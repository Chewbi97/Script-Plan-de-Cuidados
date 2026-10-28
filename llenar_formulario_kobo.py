"""
=============================================================
  BOT - PLANES DE CUIDADO SI-APS (KoboToolbox)
=============================================================
Requisitos:
  pip install selenium openpyxl webdriver-manager

Estructura esperada del Excel:
  Columna MICROTERRITORIO : Microterritorio del usuario
  Columna NOMBRE          : Nombre completo del integrante
  Columna TIPO            : Tipo de documento (CC, RC, PT, CE, TI, etc.)
  Columna CEDULA          : Número de documento del integrante

Uso:
  1. Pon el nombre de tu archivo Excel en ARCHIVO_EXCEL
  2. Ejecuta: python llenar_formulario_kobo.py

NOTA SOBRE ESTA VERSIÓN:
  Antes el bot hacía driver.get(URL_FORMULARIO) en cada usuario, lo que
  recargaba la página y generaba conflictos (igual que el popup de
  "registro sin grabar"). Ahora la página se carga UNA SOLA VEZ al
  inicio; después de cada envío exitoso el propio formulario de Kobo
  se vacía y queda listo para el siguiente registro, tal como pasa
  cuando lo llenas manualmente. El bot espera ese reseteo en vez de
  recargar la URL.
=============================================================
"""

import time
import openpyxl
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.chrome.service import Service
from webdriver_manager.chrome import ChromeDriverManager

# ─────────────────────────────────────────────
#  CONFIGURACIÓN — edita solo esta sección
# ─────────────────────────────────────────────

ARCHIVO_EXCEL = "base_datos.xlsx"
HOJA = 0
FILA_INICIO = 2

COL_MICROTERRITORIO = "MICROTERRITORIO"
COL_NOMBRE = "NOMBRE"
COL_TIPO = "TIPO"
COL_CEDULA = "CEDULA"

CODIGO_EBS = "7"
TERRITORIO = "7"
MI_CEDULA = "1091672216"

URL_FORMULARIO = "https://ee.kobotoolbox.org/x/loQfSZwv"
PAUSA = 0.8

# Cuánto esperar máximo a que el formulario se vacíe solo tras un envío
TIMEOUT_RESET_FORMULARIO = 20

BASE = "/arZq7ZMECnpaTa9DpFwh59"

# ── Names exactos sección FAMILIAR (group_hn1ol83)
N_EBS_FAM = f"{BASE}/group_hn1ol83/C_digo_de_EBS"
N_TERRITORIO_FAM = f"{BASE}/group_hn1ol83/Territorio"
N_MICRO_FAM = f"{BASE}/group_hn1ol83/Microterritorio"
N_FAMILIA_FAM = f"{BASE}/group_hn1ol83/C_digo_Familia_Hogar"

# ── Names exactos sección INDIVIDUAL (group_dx2ga91)
N_EBS_IND = f"{BASE}/group_dx2ga91/C_digo_de_EBS_001"
N_FAMILIA_IND = f"{BASE}/group_dx2ga91/C_digo_Familia_Hogar_001"
N_TERRITORIO_IND = f"{BASE}/group_dx2ga91/Territorio_001"
N_MICRO_IND = f"{BASE}/group_dx2ga91/Micro_territorio"

TIPOS_DOCUMENTO = {
    "AS": "AS. Adulto sin identificación",
    "CC": "CC. Cédula de ciudadanía",
    "CD": "CD. Carné Diplomático",
    "CE": "CE. Cédula de extranjería",
    "MS": "MS. Menor sin identificación",
    "NV": "NV. Certificado de nacido vivo",
    "PE": "PE. Permiso Especial de permanencia",
    "PT": "PT. Permiso por protección temporal",
    "RC": "RC. Registro civil",
    "TI": "TI. Tarjeta de Identidad",
}

# ─────────────────────────────────────────────
#  HELPERS
# ─────────────────────────────────────────────


def limpiar_cedula(valor):
    return str(valor).replace(".", "").replace(",", "").replace(" ", "").strip()


def escribir_por_name(driver, name, valor):
    try:
        campo = WebDriverWait(driver, 10).until(
            EC.presence_of_element_located((By.NAME, name))
        )
        driver.execute_script("arguments[0].scrollIntoView(true);", campo)
        time.sleep(0.3)
        campo.clear()
        campo.send_keys(str(valor))
        time.sleep(PAUSA)
        print(f"    ✔ {name.split('/')[-1]} = {valor}")
    except Exception as e:
        print(f"    ✘ No encontrado: {name.split('/')[-1]} → {e}")


def click_radio_por_value(driver, value):
    """Click en un radio input por su atributo value."""
    try:
        radio = driver.find_element(
            By.XPATH, f"//input[@type='radio'][@value='{value}']"
        )
        driver.execute_script("arguments[0].click();", radio)
        time.sleep(PAUSA)
        return True
    except Exception:
        return False


def click_label_contiene(driver, texto, indice=0):
    """Click en el n-ésimo label que contenga el texto."""
    labels = driver.find_elements(
        By.XPATH, f"//label[contains(normalize-space(),'{texto}')]"
    )
    if len(labels) > indice:
        driver.execute_script("arguments[0].scrollIntoView(true);", labels[indice])
        time.sleep(0.3)
        labels[indice].click()
        time.sleep(PAUSA)
        print(f"    ✔ Click label[{indice}]: {texto}")
        return True
    print(f"    ✘ No encontrado label[{indice}]: {texto}")
    return False


def click_label_exacto(driver, texto, indice=0):
    """Click en el n-ésimo label con texto exacto."""
    labels = driver.find_elements(By.XPATH, f"//label[normalize-space()='{texto}']")
    if not labels:
        labels = driver.find_elements(
            By.XPATH, f"//label[contains(normalize-space(),'{texto}')]"
        )
    if len(labels) > indice:
        driver.execute_script("arguments[0].scrollIntoView(true);", labels[indice])
        time.sleep(0.3)
        labels[indice].click()
        time.sleep(PAUSA)
        print(f"    ✔ Click label exacto[{indice}]: {texto}")
        return True
    print(f"    ✘ No encontrado label exacto[{indice}]: {texto}")
    return False


def escribir_por_label(driver, texto_label, valor):
    campos = driver.find_elements(
        By.XPATH,
        f"//*[contains(normalize-space(text()),'{texto_label}')]/following::input[1]",
    )
    if not campos:
        campos = driver.find_elements(
            By.XPATH,
            f"//*[contains(normalize-space(text()),'{texto_label}')]/following::textarea[1]",
        )
    if campos:
        driver.execute_script("arguments[0].scrollIntoView(true);", campos[0])
        campos[0].clear()
        campos[0].send_keys(str(valor))
        time.sleep(PAUSA)
        print(f"    ✔ {texto_label} = {valor}")
    else:
        print(f"    ✘ No encontrado campo: {texto_label}")


def descartar_borrador_si_aparece(driver, espera=4):
    """Si aparece el popup 'Un registro sin grabar' -> click Descartar.
    Solo hace falta revisarlo justo después de cargar la página por
    primera vez; ya no se llama entre usuarios porque la página no
    vuelve a recargarse."""
    try:
        descartar = WebDriverWait(driver, espera).until(
            EC.element_to_be_clickable(
                (By.XPATH, "//*[contains(normalize-space(),'DESCARTAR')]")
            )
        )
        descartar.click()
        time.sleep(2)
        print("    OK Draft anterior descartado")
    except Exception:
        pass


def esperar_formulario_listo(driver, timeout=TIMEOUT_RESET_FORMULARIO):
    """
    Tras un envío exitoso, Kobo/Enketo vacía el formulario solo y lo
    deja listo para el siguiente registro (sin recargar la página).
    Esta función espera a que ese reseteo realmente haya ocurrido,
    usando como referencia el campo de código de EBS de la sección
    familiar (vuelve a estar vacío cuando el formulario está listo).
    """
    try:
        WebDriverWait(driver, timeout).until(
            lambda d: (d.find_element(By.NAME, N_EBS_FAM).get_attribute("value") or "")
            == ""
        )
        return True
    except Exception as e:
        print(
            f"    ⚠ El formulario no se vació solo a tiempo ({e}). "
            f"Continuando de todas formas, revisa manualmente si hace falta."
        )
        return False


# ─────────────────────────────────────────────
#  LEER EXCEL
# ─────────────────────────────────────────────


def leer_usuarios(archivo, hoja_idx, fila_inicio):
    wb = openpyxl.load_workbook(archivo)
    ws = wb.worksheets[hoja_idx]
    encabezados = {str(cell.value).strip().upper(): cell.column for cell in ws[1]}

    def idx(col):
        key = col.upper()
        if key not in encabezados:
            raise ValueError(
                f"Columna '{col}' no encontrada. Columnas: {list(encabezados.keys())}"
            )
        return encabezados[key] - 1

    i_micro = idx(COL_MICROTERRITORIO)
    i_nombre = idx(COL_NOMBRE)
    i_tipo = idx(COL_TIPO)
    i_cedula = idx(COL_CEDULA)

    usuarios = []
    for row in ws.iter_rows(min_row=fila_inicio, values_only=True):
        nombre = row[i_nombre]
        cedula = row[i_cedula]
        if nombre and cedula:
            tipo_raw = str(row[i_tipo]).strip().upper() if row[i_tipo] else "CC"
            usuarios.append(
                {
                    "microterritorio": (
                        str(row[i_micro]).strip() if row[i_micro] else ""
                    ),
                    "nombre": str(nombre).strip(),
                    "tipo": tipo_raw,
                    "cedula": limpiar_cedula(cedula),
                }
            )
    return usuarios


# ─────────────────────────────────────────────
#  LLENAR UN FORMULARIO
# ─────────────────────────────────────────────


def llenar_formulario(driver, usuario, consecutivo):
    """
    IMPORTANTE: esta función ya NO navega a la URL ni recarga la página.
    Asume que el formulario ya está cargado y vacío (porque main() lo
    cargó una sola vez al principio, o porque el usuario anterior
    terminó con un envío exitoso que dejó el formulario reseteado).
    """
    codigo_familia = str(consecutivo).zfill(3)
    tipo_texto = TIPOS_DOCUMENTO.get(usuario["tipo"], "CC. Cédula de ciudadanía")

    print(
        f"\n→ {usuario['nombre']} | {usuario['tipo']} {usuario['cedula']} | Familia: {codigo_familia}"
    )

    # ══════════════════════════════════════════
    #  SECCIÓN 1: PLAN DE CUIDADO FAMILIAR
    # ══════════════════════════════════════════
    print("  [1] Plan Familiar")

    # Radio "Cuidado de la familia" → "-"
    radios = driver.find_elements(By.XPATH, "//input[@type='radio']")
    if radios:
        driver.execute_script("arguments[0].click();", radios[0])
        time.sleep(PAUSA)
        print("    ✔ Cuidado de la familia = -")

    escribir_por_name(driver, N_EBS_FAM, CODIGO_EBS)
    escribir_por_name(driver, N_TERRITORIO_FAM, TERRITORIO)
    escribir_por_name(driver, N_MICRO_FAM, usuario["microterritorio"])
    escribir_por_name(driver, N_FAMILIA_FAM, codigo_familia)

    # Radio "Identificación integrante del equipo de salud" → "-"
    radios = driver.find_elements(By.XPATH, "//input[@type='radio']")
    if len(radios) >= 2:
        driver.execute_script("arguments[0].click();", radios[1])
        time.sleep(PAUSA)
        print("    ✔ Identificación equipo salud = -")

    # Tipo de documento profesional → checkbox CC
    click_label_contiene(driver, "CC", indice=0)

    # Número de documento profesional (sección familiar)
    campos_numdoc = driver.find_elements(
        By.XPATH,
        "//*[contains(normalize-space(text()),'Número de documento')]/following::input[1]",
    )
    if campos_numdoc:
        driver.execute_script("arguments[0].scrollIntoView(true);", campos_numdoc[0])
        campos_numdoc[0].clear()
        campos_numdoc[0].send_keys(MI_CEDULA)
        time.sleep(PAUSA)
        print(f"    ✔ Número de documento = {MI_CEDULA}")

    # Acción: Educación en promoción...
    click_label_contiene(
        driver, "Educación en promoción y mantenimiento de la salud familiar", indice=0
    )

    # Respuesta: En sitio (primera)
    click_label_contiene(driver, "En sitio", indice=0)

    # ══════════════════════════════════════════
    #  SECCIÓN 2: PLAN DE CUIDADO INDIVIDUAL
    # ══════════════════════════════════════════
    print("  [2] Plan Individual")

    radios = driver.find_elements(By.XPATH, "//input[@type='radio']")
    if len(radios) >= 3:
        driver.execute_script("arguments[0].click();", radios[2])
        time.sleep(PAUSA)
        print("    ✔ Cuidado de la persona = -")

    escribir_por_name(driver, N_EBS_IND, CODIGO_EBS)
    escribir_por_name(driver, N_FAMILIA_IND, codigo_familia)
    escribir_por_name(driver, N_TERRITORIO_IND, TERRITORIO)
    escribir_por_name(driver, N_MICRO_IND, usuario["microterritorio"])

    # ══════════════════════════════════════════
    #  SECCIÓN 3: DATOS DEL INTEGRANTE
    # ══════════════════════════════════════════
    print(f"  [3] Integrante: {tipo_texto}")

    click_label_contiene(driver, tipo_texto, indice=0)

    campos_num = driver.find_elements(
        By.XPATH,
        "//*[contains(normalize-space(text()),'Número de identificación')]/following::input",
    )
    if campos_num:
        driver.execute_script("arguments[0].scrollIntoView(true);", campos_num[0])
        campos_num[0].clear()
        campos_num[0].send_keys(usuario["cedula"])
        time.sleep(PAUSA)
        print(f"    ✔ Número identificación integrante = {usuario['cedula']}")

    escribir_por_label(driver, "Nombres y Apellidos del integrante", usuario["nombre"])

    labels_cc = driver.find_elements(By.XPATH, "//label[normalize-space()='CC']")
    if len(labels_cc) >= 2:
        driver.execute_script("arguments[0].scrollIntoView(true);", labels_cc[1])
        labels_cc[1].click()
        time.sleep(PAUSA)
        print("    ✔ Identificación equipo salud = CC")

    campos_eq = driver.find_elements(
        By.XPATH,
        "//*[contains(normalize-space(text()),'Número de identificación integrante del equipo')]/following::input",
    )
    if campos_eq:
        driver.execute_script("arguments[0].scrollIntoView(true);", campos_eq[0])
        campos_eq[0].clear()
        campos_eq[0].send_keys(MI_CEDULA)
        time.sleep(PAUSA)
        print(f"    ✔ Número identificación equipo = {MI_CEDULA}")

    # ══════════════════════════════════════════
    #  SECCIÓN 4: AUTOCUIDADO
    # ══════════════════════════════════════════
    print("  [4] Autocuidado")
    click_label_contiene(driver, "Autocuidado", indice=0)

    # ══════════════════════════════════════════
    #  SECCIÓN 5: RESPUESTA FINAL
    # ══════════════════════════════════════════
    print("  [5] Respuesta final")

    labels_sitio = driver.find_elements(
        By.XPATH, "//label[contains(normalize-space(),'En sitio')]"
    )
    if labels_sitio:
        driver.execute_script("arguments[0].scrollIntoView(true);", labels_sitio[-1])
        labels_sitio[-1].click()
        time.sleep(PAUSA)
        print("    ✔ Respuesta = En sitio")

    labels_no = driver.find_elements(By.XPATH, "//label[normalize-space()='No']")
    if labels_no:
        driver.execute_script("arguments[0].scrollIntoView(true);", labels_no[-1])
        labels_no[-1].click()
        time.sleep(PAUSA)
        print("    ✔ Remisión = No")

    escribir_por_label(driver, "A dónde remitió", "MEDICNA GENERAL Y LECTURA DE EXAMENES")

    # ══════════════════════════════════════════
    #  SUBMIT
    # ══════════════════════════════════════════
    print("  [✓] Enviando...")
    try:
        driver.execute_script("window.scrollTo(0, document.body.scrollHeight);")
        time.sleep(2)

        btn = WebDriverWait(driver, 10).until(
            EC.element_to_be_clickable((By.ID, "submit-form"))
        )
        driver.execute_script("arguments[0].scrollIntoView(true);", btn)
        time.sleep(1)
        driver.execute_script("arguments[0].click();", btn)
        time.sleep(3)

        # Si aparece Alert de error, cerrarlo y reintentar hasta 3 veces
        for intento in range(3):
            try:
                ok_btn = driver.find_element(
                    By.XPATH,
                    "//button[normalize-space()='OK'] | //a[normalize-space()='OK']",
                )
                ok_btn.click()
                print(f"  [!] Alert cerrado, reintentando ({intento+1}/3)...")
                time.sleep(2)
                btn = WebDriverWait(driver, 10).until(
                    EC.element_to_be_clickable((By.ID, "submit-form"))
                )
                driver.execute_script("arguments[0].click();", btn)
                time.sleep(3)
            except Exception:
                break  # No hay Alert, formulario enviado

        # En vez de recargar la página, esperamos a que Kobo vacíe
        # el formulario solo y quede listo para el siguiente registro.
        esperar_formulario_listo(driver)
        print(f"  [✓] Enviado: {usuario['nombre']}")
    except Exception as e:
        print(f"  [!] Error en submit: {e}")


# ─────────────────────────────────────────────
#  MAIN
# ─────────────────────────────────────────────


def main():
    print("=" * 55)
    print("  BOT PLANES DE CUIDADO SI-APS")
    print("=" * 55)

    try:
        usuarios = leer_usuarios(ARCHIVO_EXCEL, HOJA, FILA_INICIO)
        print(f"\n✔ {len(usuarios)} usuarios cargados\n")
        for u in usuarios:
            print(
                f"  • {u['nombre']} | {u['tipo']} | {u['cedula']} | {u['microterritorio']}"
            )
        print()
    except Exception as e:
        print(f"\n✘ Error leyendo Excel: {e}")
        return

    input("Presiona ENTER para comenzar (o Ctrl+C para cancelar)...")

    options = webdriver.ChromeOptions()
    options.add_argument("--start-maximized")
    options.add_argument("--disable-notifications")

    driver = webdriver.Chrome(
        service=Service(ChromeDriverManager().install()), options=options
    )

    exitosos = []
    fallidos = []

    try:
        # La página se carga UNA SOLA VEZ, no en cada usuario.
        driver.get(URL_FORMULARIO)
        time.sleep(5)
        descartar_borrador_si_aparece(driver)

        for i, usuario in enumerate(usuarios, 1):
            print(f"\n[{i}/{len(usuarios)}]", end="")
            try:
                llenar_formulario(driver, usuario, i)
                exitosos.append(usuario["nombre"])
            except Exception as e:
                print(f"  [✘] Error: {e}")
                fallidos.append(usuario["nombre"])
                time.sleep(2)
    finally:
        driver.quit()

    print("\n" + "=" * 55)
    print("  RESUMEN FINAL")
    print("=" * 55)
    print(f"  ✔ Exitosos : {len(exitosos)}")
    print(f"  ✘ Fallidos : {len(fallidos)}")
    if fallidos:
        print("\n  Usuarios con error:")
        for u in fallidos:
            print(f"    • {u}")
    print("=" * 55)


if __name__ == "__main__":
    main()
