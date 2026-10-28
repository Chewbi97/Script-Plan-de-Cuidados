# -*- coding: utf-8 -*-
"""
=============================================================
  BOT - CARACTERIZACIÓN SI-APS RES 272 DE 2026 (ZONA RURAL)
=============================================================
Requisitos:  pip install selenium openpyxl webdriver-manager beautifulsoup4

Uso:
  python caracterizacion_rural_kobo.py                    -> carga las familias
  python caracterizacion_rural_kobo.py --validar F.html   -> prueba SIN navegador
        que cada pregunta/opción que usa el bot existe en el HTML guardado.

Excel (hoja "Caracterizacion"), columnas obligatorias (igual que antes):
  FECHA | MICROTERRITORIO | NOMBRE | TIPO | IDENTIFICACION |
  FECHA DE NACIMIENTO | TELEFONO | SEXO | EDAD | REGIMEN | EAPB
Columnas OPCIONALES (si existen y tienen dato, el bot las usa):
  DISCAPACIDAD (X) / FISICA, AUDITIVA, VISUAL... (X) / TIPO DISCAPACIDAD
  GESTANTE (X) | PESO | TALLA | CINTURA | SISTOLICA | DIASTOLICA
  NACIONALIDAD | OCUPACION | ETNIA | NIVEL EDUCATIVO | ROL
Hoja opcional "Veredas": MICROTERRITORIO | PERFIL (COMPLETO / POZO)

Regla de familia: mismo TELEFONO = misma familia. Sin teléfono, "-" o
"NO TIENE" = familia unipersonal.

REGLA DE DATOS: lo que no está en el Excel NO se inventa. Los números
clínicos no medidos van en 0 (= "no medido") y las respuestas de
vivienda salen de la tabla PERFILES / VIVIENDA_COMUN (supuestos
declarados, editables aquí abajo). Cada registro lleva una nota fija
en "Observaciones de la situación" (NOTA_OBSERVACIONES).
=============================================================
"""
import re, sys, time, csv, unicodedata
from datetime import datetime, date

# ─────────────────────────────────────────────
#  CONFIGURACIÓN
# ─────────────────────────────────────────────
ARCHIVO_EXCEL = "base_datos.xlsx"
HOJA_NOMBRE = "Caracterizacion"
HOJA_VEREDAS = "Veredas"
# El HTML guardado apunta a este link (el mismo del bot anterior). Confírmalo.
URL_FORMULARIO = "https://ee.kobotoolbox.org/x/iK2YkkfP"

DEPARTAMENTO = "Norte de Santander"
MUNICIPIO = "Ocaña"
ID_EQUIPO = "7"          # 10. Número de identificación del equipo de salud
TERRITORIO = "7"         # [Supuesto] igual que en el bot de planes de cuidado
CODIGO_FICHA = "7"
PRESTADOR = "e_s_e_hospital_emiro"     # única opción del formulario
PERFIL_PROFESIONAL = "auxiliar_de_enfermer"  # única opción del formulario
PERFIL_VEREDA_DEFECTO = "COMPLETO"

NOTA_OBSERVACIONES = ("BASE POBLACIONAL: medidas antropometricas y datos de "
                      "vivienda no verificados en visita.")

TIMEOUT_RESET = 45
PAUSA = 0.5
LOG_CSV = "log_carga_rural.csv"

# ─────────────────────────────────────────────
#  SUPUESTOS DE VIVIENDA (se aplican a todas las familias)
#  ("radio", valor) | ("check", [valores]) | ("num", n) | ("text", t)
#  Los valores pueden ser un prefijo único de la opción real.
# ─────────────────────────────────────────────
VIVIENDA_COMUN = {
    "28": ("radio", "bajo_bajo"), "33": ("radio", "no"),
    "35": ("radio", "casa"),
    "38": ("check", ["sala_comedor"]), "39": ("check", ["sala_comedor"]),
    "40": ("check", ["lavamanos"]), "41": ("radio", "luz_el_ctrica"),
    "43": ("radio", "concreto"),
    "44": ("radio", "si"), "45": ("radio", "si"), "46": ("radio", "si"),
    "47": ("check", ["motocicleta"]),
    "49": ("check", ["distancias_a_recorrer"]),
    "50": ("radio", "no"), "51": ("radio", "no"), "52": ("radio", "no"),
    "62": ("radio", "no"), "63": ("radio", "no"), "65": ("radio", "no"),
    "67": ("check", ["en_conjunto_con_los_residuos_ordinarios_"]),
    "68": ("radio", "no"),
    "69": ("check", ["separaci_n_de_residuos_no_aprovechables_"]),
    "70": ("radio", "retiro_de_polvo"),
    "71": ("check", ["gas_licuado_del_petr_leo__gas_propano"]),
    "73": ("check", ["barrido_de_pisos_con_levantamiento_de_po"]),
    "74": ("check", ["el_ctrica__de_gas__petr_leo__gasolina__k"]),
    "75": ("radio", "no"), "76": ("check", ["moscas"]),
    "77": ("check", ["manejo_higi_nico_de_los_residuos_s_lidos"]),
    "79": ("check", ["orden_e_higiene_en_las_viviendas_y_el_pe"]),
    "81": ("num", 0), "82": ("num", 0), "83": ("num", 0), "84": ("num", 0),
    "85": ("check", ["compa_a"]), "86": ("check", ["totalmente_confinados"]),
    "87": ("radio", "no"), "88": ("radio", "no"),
    "89": ("radio", "no"), "90": ("radio", "no"),
    "91": ("check", ["manejo_higi_nico_de_residuos_s_lidos"]),
    "92": ("radio", "tiendas_de_barrio"),
    "93": ("check", ["en_conjunto_con_los_residuos_ordinarios_"]),
    "94": ("radio", "si"), "95": ("radio", "si"), "96": ("radio", "si"),
    # familia
    "100": ("radio", "7_a_10"), "101": ("radio", "no"),
    "102": ("radio", "no_aplica"),
    "104": ("check", ["la_comunicaci_n_en_la_familia_est__basad"]),
    "105": ("radio", "cuenta_con_redes_de_apoyo_sociales_prote"),
    "106": ("check", ["facilita_la_circulaci_n_del_aire_en_la_v"]),
    "107": ("check", ["lavado_frecuente_de_manos"]), "108": ("radio", "no"),
}
# Saneamiento básico según la vereda (hoja "Veredas")
PERFILES = {
    "COMPLETO": {
        "53": ("radio", "acueducto_administrado_por_empresa_prest"),
        "54": ("radio", "24_horas"),
        "55": ("radio", "no_tiene_tanque_de_almacenamiento"),
        "56": ("radio", "no_aplica"), "57": ("radio", "m_s_de_10_metros"),
        "58": ("check", ["sanitario_conectado_a_alcantarillado"]),
        "59": ("check", ["alcantarillado"]),
        "60": ("check", ["recipientes_con_tapa"]),
        "61": ("radio", "recolecci_n_por_parte_del_servicio_de_as"),
    },
    "POZO": {
        "53": ("radio", "acueducto_veredal_o_comunitario"),
        "54": ("radio", "de_12_a_23_horas"),
        "55": ("radio", "no_tiene_tanque_de_almacenamiento"),
        "56": ("radio", "no_aplica"), "57": ("radio", "m_s_de_10_metros"),
        "58": ("check", ["sanitario_conectado_a_pozo_s_ptico"]),
        "59": ("check", ["pozo_s_ptico"]),
        "60": ("check", ["recipientes_con_tapa"]),
        "61": ("radio", "recolecci_n_por_parte_del_servicio_de_as"),
    },
}
# Preguntas obligatorias SIN opción neutra que solo se responden si el
# formulario las muestra. None = el bot se detiene y te la señala.
RESPUESTAS_FORZADAS = {
    "149": ("radio", "peso_adecuado_para_la_talla_o_imc_adecua"),
    "154": None,
    "155": ("radio", "no_acudi__a_ning_n_servicio_de_salud"),
    "194": ("radio", "ninguno___no_fue_ofertado"),
}
NEUTROS = ["no_aplica", "no_registra", "ninguna", "ninguno", "sin_discapacidad"]

TIPOS_DOC = {  # prefijos de las opciones de la pregunta 113
    "AS": "as__adulto", "RC": "rc__registro", "TI": "ti__tarjeta",
    "CC": "cc__cedula", "CD": "cd__carn", "CE": "ce__c_dula",
    "MS": "ms__menor", "NV": "nv__certificado", "PT": "pt__permiso",
}
REGIMEN_VAL = {
    "SUB": "subsidiado", "SUBSIDIADO": "subsidiado",
    "CON": "contributivo", "CONTRIBUTIVO": "contributivo",
    "ESPECIAL": "especial", "EXCEPCION": "excepci_n",
    "NO AFILIADO": "no_afiliado",
}
TIPOS_DISC = {"FISICA": "f_sica", "AUDITIVA": "auditiva", "VISUAL": "visual",
              "SORDOCEGUERA": "sordoceguera", "INTELECTUAL": "intelectual",
              "PSICOSOCIAL": "psicosocial__mental", "MENTAL": "psicosocial__mental",
              "MULTIPLE": "m_ltiple"}
ETNIAS = {"INDIGENA": "ind_gena", "ROM": "rom__gitanos", "GITANO": "rom__gitanos",
          "NEGRO": "negro", "AFRO": "afrocolombiano", "RAIZAL": "raizal",
          "PALENQUERO": "palenquero", "NINGUNA": "ninguna"}


class CampoError(Exception):
    """Detiene TODO el proceso (no sigue con la siguiente familia)."""


class DatoFaltante(CampoError):
    """Falta un dato que el bot se niega a inventar."""


# ─────────────────────────────────────────────
#  UTILIDADES
# ─────────────────────────────────────────────
def norm(t):
    t = unicodedata.normalize("NFKD", str(t or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", t).strip().upper()


def limpiar(v):
    return re.sub(r"[.,\s]", "", str(v or "")).strip()


def tel_valido(t):
    return norm(t) not in ("", "-", "NO TIENE", "NONE", "NAN", "0")


def es_x(v):
    return norm(v) in ("X", "SI", "1", "TRUE")


def a_fecha(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()[:10]
    for f in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def edad_calc(fn, hoy=None):
    hoy = hoy or date.today()
    return hoy.year - fn.year - ((hoy.month, hoy.day) < (fn.month, fn.day))


def num_o_none(v):
    try:
        return float(str(v).replace(",", ".")) if str(v or "").strip() else None
    except ValueError:
        return None


def dividir_nombre(n):
    p = n.strip().split()
    if len(p) == 1:
        return p[0], "", p[0], "NO REGISTRA"
    if len(p) == 2:
        return p[0], "", p[1], "NO REGISTRA"
    if len(p) == 3:
        return p[0], "", p[1], p[2]
    return p[0], p[1], p[2], " ".join(p[3:])


def clasif_imc_adulto(imc):
    if imc < 18.5:
        return "delgadez"
    if imc < 25:
        return "peso_adecuado_para_la_talla_o_imc_adecua"
    if imc < 30:
        return "sobrepeso"
    return "obesidad"


def clasif_ta(s, d):
    if s > 180 or d > 120:
        return "crisis_hipertensiva"
    if s >= 140 or d >= 90:
        return "alta___hipertensi_n_nivel_2"
    if s >= 130 or d >= 80:
        return "alta___hipertensi_n_nivel_1"
    if s >= 120:
        return "elevada__sist_lica"
    return "normal__sist_lica"


# ─────────────────────────────────────────────
#  LEER EXCEL
# ─────────────────────────────────────────────
OPCIONALES = {
    "disc_x": ["DISCAPACIDAD"], "tipo_disc": ["TIPO DISCAPACIDAD"],
    "gestante": ["GESTANTE", "EMBARAZ"], "peso": ["PESO"], "talla": ["TALLA"],
    "cintura": ["CINTURA"], "sist": ["SISTOLICA"], "diast": ["DIASTOLICA"],
    "nacionalidad": ["NACIONALIDAD"], "ocupacion": ["OCUPACION"],
    "etnia": ["ETNIA"], "nivel": ["NIVEL EDUCATIVO", "ESCOLARIDAD"],
    "rol": ["ROL"],
}


def _buscar(hdr, nombres, usados):
    for n in nombres:                      # 1) exacto
        for i, h in enumerate(hdr):
            if h == n and i not in usados:
                return i
    for n in nombres:                      # 2) contiene
        for i, h in enumerate(hdr):
            if n in h and i not in usados:
                return i
    return None


def leer_veredas(wb):
    if HOJA_VEREDAS not in wb.sheetnames:
        return {}
    ws = wb[HOJA_VEREDAS]
    out = {}
    for r in list(ws.iter_rows(values_only=True))[1:]:
        if r and r[0] and len(r) > 1 and r[1]:
            out[norm(r[0])] = norm(r[1])
    return out


def leer_familias(archivo, hoja):
    import openpyxl
    wb = openpyxl.load_workbook(archivo, read_only=True, data_only=True)
    ws = wb[hoja]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [norm(c) for c in rows[0]]
    usados = set()

    def obl(*nombres):
        i = _buscar(hdr, nombres, usados)
        if i is None:
            raise ValueError(f"No encontré la columna '{nombres[0]}'. Columnas: {hdr}")
        usados.add(i)
        return i

    iF, iM = obl("FECHA"), obl("MICROTERRITORIO")
    iN, iT = obl("NOMBRE"), obl("TIPO")
    iI, iNac = obl("IDENTIFICACION"), obl("NACIMIENTO")
    iTel, iS = obl("TELEFONO"), obl("SEXO")
    iE, iR, iEapb = obl("EDAD"), obl("REGIMEN"), obl("EAPB")

    opc = {}
    for k, nombres in OPCIONALES.items():
        i = _buscar(hdr, nombres, usados)
        if i is not None:
            opc[k] = i
            usados.add(i)
    # columnas por tipo de discapacidad (FISICA, AUDITIVA, ...)
    cols_tipo = {}
    for i, h in enumerate(hdr):
        if i not in usados and h in TIPOS_DISC:
            cols_tipo[TIPOS_DISC[h]] = i
    print("  Columnas opcionales detectadas:",
          {k: rows[0][v] for k, v in opc.items()} or "ninguna")
    if cols_tipo:
        print("  Columnas de tipo de discapacidad:", sorted(cols_tipo))
    faltan = [k for k in OPCIONALES if k not in opc]
    print("  Sin columna (se usa la regla por defecto):", faltan)

    def g(row, k):
        return row[opc[k]] if k in opc and opc[k] < len(row) else None

    personas = []
    for row in rows[1:]:
        if not row or not row[iN]:
            continue
        tel_raw = row[iTel]
        fn = a_fecha(row[iNac])
        m_e = re.search(r"\d+", str(row[iE] if row[iE] is not None else ""))
        edad = int(m_e.group()) if m_e else (edad_calc(fn) if fn else None)
        discs = [v for v, i in cols_tipo.items() if es_x(row[i])]
        personas.append({
            "fecha": a_fecha(row[iF]), "micro": str(row[iM] or "").strip(),
            "nombre": str(row[iN]).strip(), "tipo": norm(row[iT]) or "CC",
            "id": limpiar(row[iI]), "fecha_nac": fn, "edad": edad,
            "telefono": limpiar(tel_raw) if tel_valido(tel_raw) else None,
            "sexo": norm(row[iS]), "regimen": norm(row[iR]),
            "eapb": str(row[iEapb] or "").strip(),
            "gestante": es_x(g(row, "gestante")),
            "disc_x": es_x(g(row, "disc_x")),
            "tipo_disc": norm(g(row, "tipo_disc")), "discs": discs,
            "peso": num_o_none(g(row, "peso")), "talla": num_o_none(g(row, "talla")),
            "cintura": num_o_none(g(row, "cintura")),
            "sist": num_o_none(g(row, "sist")), "diast": num_o_none(g(row, "diast")),
            "nacionalidad": str(g(row, "nacionalidad") or "").strip(),
            "ocupacion": str(g(row, "ocupacion") or "").strip(),
            "etnia": norm(g(row, "etnia")), "nivel": norm(g(row, "nivel")),
            "rol": norm(g(row, "rol")),
        })
    familias, idx_tel = [], {}
    for p in personas:
        t = p["telefono"]
        if t and t in idx_tel:
            familias[idx_tel[t]].append(p)
        elif t:
            idx_tel[t] = len(familias)
            familias.append([p])
        else:
            familias.append([p])
    return familias, leer_veredas(wb)


# ─────────────────────────────────────────────
#  REGLAS DE RESPUESTA  (devuelven listas (pregunta, tipo, valor))
# ─────────────────────────────────────────────
def respuestas_hogar(familia, cod, perfil, run):
    p0, n = familia[0], len(familia)
    micro = p0["micro"]
    apell = p0["nombre"].split()
    alias = "FAMILIA " + (apell[-2] if len(apell) >= 3 else apell[-1])
    hab = max(1, round(n / 2 + 0.5))
    adultos = sum(1 for p in familia if (p["edad"] or 0) >= 18)
    tel = next((p["telefono"] for p in familia if p["telefono"]), None)
    f_dil = p0["fecha"] or date.today()
    R = [("1", "radio", "si"), ("2.1", "radio", "no_aplica"),
         ("3", "text", NOTA_OBSERVACIONES), ("4", "text", DEPARTAMENTO),
         ("5", "text", MUNICIPIO), ("6", "radio", "rural"),
         ("territorio", "text", run["territorio"]),
         ("microterritorio", "text", micro), ("corregimiento", "text", micro),
         ("10", "num", run["equipo"]), ("11", "radio", PRESTADOR),
         ("12", "radio", run["tipo_id_prof"]), ("13", "num", run["cedula"]),
         ("13.1", "text", run["nombre_prof"]),
         ("14", "radio", PERFIL_PROFESIONAL), ("15", "num", run["ficha"]),
         ("16", "date", f_dil.strftime("%Y-%m-%d")), ("17", "radio", "hogar"),
         ("18", "text", alias), ("19", "text", p0["nombre"]),
         ("20", "radio", "no"), ("21", "text", micro), ("24", "text", micro),
         ("26", "num", int(cod)), ("27", "num", int(cod)),
         ("29", "num", 1), ("30", "num", n), ("31", "num", hab),
         ("32", "num", max(1, round(n / hab))), ("34", "num", n),
         ("97", "num", n), ("98", "text", alias),
         ("99", "radio", "unipersonal" if n == 1 else
          ("nuclear_biparental" if adultos >= 2 else "nuclear_monoparental"))]
    if tel:
        R.append(("25", "num", tel))
    tabla = dict(VIVIENDA_COMUN)
    tabla.update(PERFILES[perfil])
    for q, (t, v) in tabla.items():
        R.append((q, t, v))
    return R


def _rol(p, pos):
    if p["rol"]:
        for k, v in (("RESPONSABLE", "responsable_econ_mico_de_la_familia"),
                     ("CONYUGE", "c_nyuge_o_compa_ero_a"), ("HIJ", "hijo_a"),
                     ("HERMAN", "hermano_a"), ("PADRE", "padre_o_madre"),
                     ("MADRE", "padre_o_madre")):
            if k in p["rol"]:
                return v
        return "otro"
    if pos == 0:
        return "responsable_econ_mico_de_la_familia"
    return "hijo_a" if (p["edad"] or 0) < 18 else "otro"


def respuestas_integrante(p, pos, avisos):
    nom = p["nombre"]
    if p["edad"] is None:
        raise DatoFaltante(f"{nom}: sin EDAD ni fecha de nacimiento válida")
    if not p["fecha_nac"]:
        raise DatoFaltante(f"{nom}: fecha de nacimiento inválida o vacía")
    if not p["eapb"]:
        raise DatoFaltante(f"{nom}: EAPB vacía (obligatoria en el formulario)")
    if not p["id"]:
        raise DatoFaltante(f"{nom}: IDENTIFICACION vacía")
    edad, tipo = p["edad"], p["tipo"]
    if tipo not in TIPOS_DOC:
        raise DatoFaltante(f"{nom}: tipo de documento '{tipo}' no existe en el formulario rural")
    if p["regimen"] not in REGIMEN_VAL:
        raise DatoFaltante(f"{nom}: régimen '{p['regimen']}' no reconocido")
    fem = p["sexo"].startswith("F")
    if not (fem or p["sexo"].startswith("M")):
        raise DatoFaltante(f"{nom}: SEXO debe ser FEMENINO o MASCULINO")
    n1, n2, a1, a2 = dividir_nombre(nom)
    # nacionalidad: solo se deduce cuando es un hecho
    if p["nacionalidad"]:
        nac = p["nacionalidad"]
    elif tipo in ("CC", "TI", "RC", "NV"):
        nac = "Colombiana"
    elif tipo == "PT":
        nac = "Venezolana"           # el PPT solo se otorga a venezolanos
    else:
        raise DatoFaltante(f"{nom}: documento {tipo} requiere columna NACIONALIDAD")
    estatus = "regular" if tipo in ("PT", "CE", "CD") else "no_aplica"
    # discapacidad
    discs = list(p["discs"])
    if p["disc_x"]:
        for k, v in TIPOS_DISC.items():
            if k in p["tipo_disc"] and v not in discs:
                discs.append(v)
        if not discs:
            raise DatoFaltante(f"{nom}: marcado con X en discapacidad pero falta el tipo "
                               f"(columna TIPO DISCAPACIDAD o columnas por tipo)")
    # etnia
    et = "ninguna"
    if p["etnia"]:
        et = next((v for k, v in ETNIAS.items() if k in p["etnia"]), None)
        if et is None:
            raise DatoFaltante(f"{nom}: ETNIA '{p['etnia']}' no reconocida")
    # nivel educativo
    if p["nivel"]:
        nivel = {"NINGUNO": "ninguno", "PREESCOLAR": "preescolar", "PRIMARIA": "b_sica_primaria",
                 "SECUNDARIA": "b_sica_secundaria", "MEDIA": "media_acad_mica_o_cl_sica",
                 "BACHILLER": "media_acad_mica_o_cl_sica"}
        nivel = next((v for k, v in nivel.items() if k in p["nivel"]), None)
        if nivel is None:
            raise DatoFaltante(f"{nom}: NIVEL EDUCATIVO '{p['nivel']}' no reconocido")
    else:
        nivel = ("preescolar" if edad < 6 else "b_sica_primaria" if edad < 12 else
                 "b_sica_secundaria" if edad < 18 else "media_acad_mica_o_cl_sica")
        avisos.add("Nivel educativo derivado de la edad cuando no hay columna")
    # protección especial (una sola opción)
    prot = ("gestante" if p["gestante"] else "persona_con_condici_n_de_discapaci" if discs
            else "ni_as__ni_os_y_adolescentes" if edad < 18 else
            "persona_adulto_mayor" if edad >= 60 else "ninguna")
    fertil = fem and 10 <= edad <= 49
    peso, talla = p["peso"] or 0, p["talla"] or 0
    imc = round(peso / ((talla / 100) ** 2), 1) if peso and talla else 0
    if imc and edad >= 18:
        clasif = clasif_imc_adulto(imc)
    else:
        clasif = None
        avisos.add("Clasificación nutricional (149) en valor por defecto: sin medidas válidas de adulto")
    R = [("109", "text", n1), ("111", "text", a1), ("112", "text", a2),
         ("113", "radio", TIPOS_DOC[tipo]), ("114", "num", p["id"]),
         ("115", "date", p["fecha_nac"].strftime("%Y-%m-%d")),
         ("116", "text", nac), ("117", "radio", estatus),
         ("118", "radio", "mujer" if fem else "hombre"),
         ("119", "radio", "femenino" if fem else "masculino"),
         ("120", "radio", "femenino" if fem else "masculino"),
         ("121", "radio", "no_responde"),
         ("124", "radio", _rol(p, pos)),
         ("125", "text", p["ocupacion"] or "NO REGISTRA"),
         ("126", "radio", nivel), ("127", "radio", REGIMEN_VAL[p["regimen"]]),
         ("128", "text", p["eapb"]), ("129", "radio", prot),
         ("130", "radio", "no_aplica"), ("131", "radio", et),
         ("132", "text", "NO APLICA" if et == "ninguna" else "NO REGISTRA"),
         ("137", "check", discs or ["sin_discapacidad"]),
         ("138", "radio", "no" if discs else "no_aplica"),
         ("139", "radio", "no"), ("140", "radio", "si" if p["gestante"] else "no"),
         ("143", "radio", "no_aplica"), ("144", "num", int(round(peso))),
         ("145", "num", int(round(p["cintura"] or 0))),
         ("146", "num", int(round(talla))), ("147", "num", int(round(imc))),
         ("148", "radio", "no_aplica"), ("150", "radio", "si"),
         ("159", "radio", "no"), ("160", "radio", "no"), ("161", "radio", "no"),
         ("173", "radio", "si" if fertil else ("no" if fem else "no_aplica")),
         ("174", "radio", "si" if p["gestante"] else ("no" if fertil else "no_aplica"))]
    if clasif:
        R.append(("149", "radio", clasif))
    if p["telefono"]:
        R.append(("122", "num", p["telefono"]))
    if p["sist"] and p["diast"]:
        R += [("sist_lica", "num", int(p["sist"])), ("diast_lica", "num", int(p["diast"])),
              ("152", "radio", clasif_ta(p["sist"], p["diast"]))]
    else:
        R.append(("152", "radio", "no_aplica"))
    if p["gestante"]:
        R.append(("180", "radio", "riesgo_no_evaluado"))
        avisos.add("Gestante: ruta materno-perinatal completar en Kubapp")
    if discs:
        avisos.add("Discapacidad: certificación RLCPD (138) marcada 'no' por defecto")
    return R


# ─────────────────────────────────────────────
#  MAPA pregunta -> data-ref (se lee de la página, no está fijo en el código)
# ─────────────────────────────────────────────
def clave_de_ref(ref):
    nombre = ref.rsplit("/", 1)[-1]
    m = re.match(r"^_(\d+(?:_\d+)*)_[A-Za-z]", nombre)
    if m:
        return m.group(1).replace("_", ".")
    for pref, k in (("Territorio", "territorio"), ("Microterritorio", "microterritorio"),
                    ("Corregimiento", "corregimiento"), ("Sist_lica", "sist_lica"),
                    ("Diast_lica", "diast_lica")):
        if nombre.startswith(pref):
            return k
    return None


def usadas_en_codigo():
    """Todas las preguntas que el bot toca explícitamente (para validar)."""
    p = {"fecha": date(2026, 1, 1), "micro": "X", "nombre": "A B C D", "tipo": "CC",
         "id": "1", "fecha_nac": date(2000, 1, 1), "edad": 26, "telefono": "3000000000",
         "sexo": "FEMENINO", "regimen": "SUB", "eapb": "E", "gestante": True,
         "disc_x": True, "tipo_disc": "FISICA", "discs": [], "peso": 60, "talla": 160,
         "cintura": 80, "sist": 110, "diast": 70, "nacionalidad": "", "ocupacion": "",
         "etnia": "", "nivel": "", "rol": ""}
    run = {"territorio": "7", "equipo": "7", "tipo_id_prof": "cc__c_dula_de_cuidadan_a",
           "cedula": "1", "nombre_prof": "X", "ficha": "7"}
    out = []
    for perfil in PERFILES:
        out += respuestas_hogar([p], "001", perfil, run)
    for var in ({}, {"gestante": False, "disc_x": False, "peso": None, "talla": None,
                     "sist": None, "diast": None, "tipo": "PT", "sexo": "MASCULINO",
                     "edad": 70, "etnia": "INDIGENA", "telefono": None, "nivel": "PRIMARIA"},
                {"edad": 5, "tipo": "RC", "disc_x": False, "gestante": False}):
        q = dict(p); q.update(var)
        out += respuestas_integrante(q, 0, set())
    return out


def validar_html(ruta):
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(open(ruta, encoding="utf-8").read(), "html.parser")
    mapa, vals = {}, {}
    for e in soup.find_all(["input", "select", "textarea"]):
        ref = e.get("data-ref") or e.get("name")
        if not ref:
            continue
        k = clave_de_ref(ref)
        if k:
            mapa.setdefault(k, ref)
        if e.get("type") in ("radio", "checkbox") and e.get("value"):
            vals.setdefault(ref, []).append(e["value"])
    errores, vistos = [], set()
    for q, tipo, v in usadas_en_codigo():
        if q not in mapa:
            errores.append(f"pregunta {q}: no existe en el HTML")
            continue
        if tipo in ("radio", "check"):
            for x in (v if isinstance(v, list) else [v]):
                if (q, x) in vistos:
                    continue
                vistos.add((q, x))
                op = vals.get(mapa[q], [])
                if x in op:
                    continue
                c = [o for o in op if o.startswith(x)]
                if len(c) != 1:
                    errores.append(f"pregunta {q}: opción '{x}' -> "
                                   f"{'ambigua ' + str(c) if c else 'no existe'}")
    for q, rf in RESPUESTAS_FORZADAS.items():
        if rf and q in mapa:
            x = rf[1]; op = vals.get(mapa[q], [])
            if x not in op and len([o for o in op if o.startswith(x)]) != 1:
                errores.append(f"forzada {q}: opción '{x}' inválida")
    print(f"Preguntas distintas verificadas: {len({q for q, _, _ in usadas_en_codigo()})}")
    if errores:
        print("ERRORES:"); [print("  -", e) for e in sorted(set(errores))]
        return False
    print("OK: todas las preguntas y opciones usadas existen en el HTML.")
    return True


# ─────────────────────────────────────────────
#  SELENIUM
# ─────────────────────────────────────────────
JS_ESTADO = """
const root = arguments[0] || document;
const enRepeat = root !== document;
const out = [];
root.querySelectorAll('label.question, fieldset.question').forEach(q => {
  if (!enRepeat && q.closest('section.or-repeat')) return;
  const els = [...q.querySelectorAll('input,select,textarea')].filter(e => e.dataset.ref || e.name);
  if (!els.length) return;
  const e0 = els[0], t = e0.type;
  const req = els.some(e => e.dataset.required === 'true()');
  let filled;
  if (t === 'radio' || t === 'checkbox') filled = els.some(e => e.checked);
  else filled = els.some(e => (e.value || '') !== '');
  const lab = q.querySelector('.question-label');
  out.push({ref: e0.dataset.ref || e0.name, kind: t, req: req, filled: filled,
            visible: q.getClientRects().length > 0,
            label: lab ? lab.innerText.trim().slice(0, 70) : '',
            vals: els.filter(e => t === 'radio' || t === 'checkbox').map(e => e.value)});
});
return out;
"""


class Bot:
    def __init__(self, driver):
        from selenium.webdriver.common.by import By
        from selenium.webdriver.support.ui import WebDriverWait
        from selenium.webdriver.support import expected_conditions as EC
        self.d, self.By, self.WW, self.EC = driver, By, WebDriverWait, EC
        self.mapa = {}

    def W(self, t=12):
        return self.WW(self.d, t)

    def leer_mapa(self):
        refs = self.d.execute_script(
            "return [...document.querySelectorAll('input,select,textarea')]"
            ".map(e => e.dataset.ref || e.name).filter(Boolean);")
        self.mapa = {}
        for r in refs:
            k = clave_de_ref(r)
            if k:
                self.mapa.setdefault(k, r)
        return self.mapa

    def ref(self, q):
        if q not in self.mapa:
            raise CampoError(f"La pregunta {q} no existe en la página cargada "
                             f"(¿cambió el formulario?)")
        return self.mapa[q]

    def _base(self, ctx):
        return ctx or self.d

    def _inputs(self, q, tipo, valor, ctx):
        ref, css = self.ref(q), self.By.CSS_SELECTOR
        b = self._base(ctx)
        sel = f'input[type="{tipo}"][data-ref="{ref}"]'
        ex = b.find_elements(css, f'{sel}[value="{valor}"]')
        if ex:
            return ex[0]
        pre = b.find_elements(css, f'{sel}[value^="{valor}"]')
        if len(pre) == 1:
            return pre[0]
        raise CampoError(f"pregunta {q}: opción '{valor}' "
                         f"{'ambigua' if pre else 'no encontrada'}")

    def _visible(self, q, ctx):
        css = self.By.CSS_SELECTOR
        els = self._base(ctx).find_elements(css, f'[data-ref="{self.ref(q)}"]')
        if not els:
            return False
        return bool(self.d.execute_script(
            "const q=arguments[0].closest('label.question, fieldset.question');"
            "return !!q && q.getClientRects().length>0;", els[0]))

    def marcar(self, q, tipo, valor, ctx=None, intentos=3):
        el = self._inputs(q, tipo, valor, ctx)
        for _ in range(intentos):
            self.d.execute_script("arguments[0].scrollIntoView({block:'center'});", el)
            if el.is_selected():
                return True
            self.d.execute_script(
                "const e=arguments[0]; e.checked=true;"
                "e.dispatchEvent(new Event('click',{bubbles:true}));"
                "e.dispatchEvent(new Event('change',{bubbles:true}));", el)
            time.sleep(PAUSA)
            if el.is_selected():
                return True
            try:
                el.click(); time.sleep(PAUSA)
            except Exception:
                pass
            if el.is_selected():
                return True
        raise CampoError(f"pregunta {q}: no se pudo marcar '{valor}'")

    def escribir(self, q, valor, ctx=None):
        css = self.By.CSS_SELECTOR
        el = self._base(ctx).find_element(css, f'[data-ref="{self.ref(q)}"]')
        self.d.execute_script("arguments[0].scrollIntoView({block:'center'});", el)
        el.clear(); el.send_keys(str(valor)); time.sleep(PAUSA)

    def fecha(self, q, valor, ctx=None):
        css = self.By.CSS_SELECTOR
        hid = self._base(ctx).find_element(css, f'[data-ref="{self.ref(q)}"]')
        lab = hid.find_element(self.By.XPATH, "./ancestor::label[1]")
        inp = lab.find_element(css, "div.widget.date input.input-small")
        self.d.execute_script("arguments[0].scrollIntoView({block:'center'});", inp)
        inp.clear(); inp.send_keys(valor)
        self.d.execute_script(
            "arguments[0].dispatchEvent(new Event('change',{bubbles:true}));", inp)
        time.sleep(PAUSA)

    def aplicar(self, respuestas, ctx=None):
        """Orden de la lista; las preguntas no visibles (irrelevantes) se omiten."""
        for q, tipo, v in respuestas:
            if not self._visible(q, ctx):
                continue
            try:
                if tipo == "radio":
                    self.marcar(q, "radio", v, ctx)
                elif tipo == "check":
                    for x in v:
                        self.marcar(q, "checkbox", x, ctx)
                elif tipo == "date":
                    self.fecha(q, v, ctx)
                else:
                    self.escribir(q, v, ctx)
            except CampoError:
                raise
            except Exception as e:
                raise CampoError(f"pregunta {q} = {v}: {e}")

    def barrido(self, ctx=None):
        """Obligatorias visibles sin responder: neutral si existe; si no, forzada;
        si no, se devuelve la lista de pendientes."""
        pend = []
        for it in self.d.execute_script(JS_ESTADO, ctx):
            if not (it["req"] and it["visible"] and not it["filled"]):
                continue
            q = clave_de_ref(it["ref"])
            tipo = "radio" if it["kind"] == "radio" else "checkbox"
            if it["kind"] in ("radio", "checkbox"):
                val = next((v for n in NEUTROS for v in it["vals"] if v.startswith(n)), None)
                if val:
                    self.marcar(q, tipo, val, ctx)
                    continue
                if q in RESPUESTAS_FORZADAS and RESPUESTAS_FORZADAS[q]:
                    self.marcar(q, tipo, RESPUESTAS_FORZADAS[q][1], ctx)
                    continue
            pend.append(f"{q}. {it['label']}")
        return pend

    # repeats de integrantes
    def _sel_rep(self):
        return f'section.or-repeat[name="{self.ref("109").rsplit("/", 1)[0]}"]'

    def contar_repeats(self):
        return len(self.d.find_elements(self.By.CSS_SELECTOR, self._sel_rep()))

    def get_repeat(self, i, t=8):
        try:
            self.W(t).until(lambda d: self.contar_repeats() > i)
        except Exception:
            pass
        reps = self.d.find_elements(self.By.CSS_SELECTOR, self._sel_rep())
        if i < len(reps):
            return reps[i]
        raise CampoError(f"No se encontró el integrante #{i + 1} (hay {len(reps)})")

    def cerrar_popups(self):
        self.d.execute_script(
            "document.querySelectorAll('.vex, .vex-overlay').forEach(e => e.remove());")

    def add_repeat(self, intentos=4):
        css, grupo = self.By.CSS_SELECTOR, self.ref("109").rsplit("/", 1)[0]
        antes, err = self.contar_repeats(), None
        for _ in range(intentos):
            try:
                self.cerrar_popups()
                btn = self.W(10).until(self.EC.element_to_be_clickable(
                    (css, f'div.or-repeat-info[data-name="{grupo}"] button.add-repeat-btn')))
                self.d.execute_script("arguments[0].scrollIntoView({block:'center'});", btn)
                try:
                    btn.click()
                except Exception:
                    self.d.execute_script("arguments[0].click();", btn)
                self.W(6).until(lambda d: self.contar_repeats() > antes)
                time.sleep(0.5)
                return
            except Exception as e:
                err = e
                time.sleep(0.7)
        raise CampoError(f"No se pudo agregar integrante: {err}")

    def descartar_borrador(self, t=5):
        try:
            b = self.W(t).until(self.EC.element_to_be_clickable(
                (self.By.XPATH, "//*[contains(normalize-space(),'DESCARTAR')]")))
            b.click(); time.sleep(2)
        except Exception:
            pass

    def formulario_vacio(self, timeout=TIMEOUT_RESET):
        css = self.By.CSS_SELECTOR
        try:
            self.W(timeout).until(lambda d: (d.find_element(
                css, f'[data-ref="{self.ref("4")}"]').get_attribute("value") or "") == "")
            time.sleep(3)
            return True
        except Exception:
            return False

    def enviar(self):
        css = self.By.CSS_SELECTOR
        self.cerrar_popups()
        self.d.execute_script("window.scrollTo(0, document.body.scrollHeight);")
        time.sleep(2)
        btn = self.W(15).until(self.EC.element_to_be_clickable((self.By.ID, "submit-form")))
        self.d.execute_script("arguments[0].click();", btn)
        time.sleep(3)
        for _ in range(3):
            try:
                ok = self.d.find_element(self.By.XPATH,
                    "//button[normalize-space()='OK'] | //a[normalize-space()='OK']")
                ok.click(); time.sleep(2)
                btn = self.W(10).until(self.EC.element_to_be_clickable((self.By.ID, "submit-form")))
                self.d.execute_script("arguments[0].click();", btn); time.sleep(3)
            except Exception:
                break
        if not self.formulario_vacio():
            falt = self.d.execute_script(
                "return [...document.querySelectorAll('.invalid-required,.invalid-constraint')]"
                ".map(q => (q.querySelector('.question-label')||{}).innerText).filter(Boolean)"
                ".slice(0, 10);")
            raise CampoError("El formulario no se envió. Preguntas con error: "
                             + (" | ".join(falt) if falt else "(no detectadas, revisa la pantalla)"))


# ─────────────────────────────────────────────
#  LLENAR UNA FAMILIA
# ─────────────────────────────────────────────
def llenar(bot, familia, consec, perfil, run, avisos):
    cod = str(consec).zfill(3)
    print(f"\n{'=' * 55}\n  FAMILIA #{cod} | {familia[0]['nombre']} | "
          f"{len(familia)} integrante(s) | perfil {perfil}\n{'=' * 55}")
    # Primero se valida TODO el Excel de la familia: si falta un dato, ni se toca el formulario
    hogar = respuestas_hogar(familia, cod, perfil, run)
    integ = [respuestas_integrante(p, i, avisos) for i, p in enumerate(familia)]
    bot.aplicar(hogar)
    for i, (p, r) in enumerate(zip(familia, integ)):
        if i > 0:
            print(f"  → Agregando integrante {i + 1}...")
            bot.add_repeat()
        print(f"  [S5.{i + 1}] {p['nombre']}")
        ctx = bot.get_repeat(i)
        bot.aplicar(r, ctx)
        pend = bot.barrido(ctx)
        if pend:
            raise CampoError(f"Integrante {p['nombre']}: obligatorias sin respuesta posible: "
                             + " | ".join(pend))
    pend = bot.barrido(None)
    if pend:
        raise CampoError("Obligatorias sin respuesta posible: " + " | ".join(pend))
    print("[✓] Enviando...")
    bot.enviar()
    print(f"  ✔ Familia #{cod} enviada")


# ─────────────────────────────────────────────
#  MAIN
# ─────────────────────────────────────────────
def pedir(msg, defecto):
    r = input(f"{msg} [{defecto}]: ").strip()
    return r or defecto


def pedir_datos_corrida():
    print("\nDatos del profesional que realiza la caracterización:")
    nombre = ""
    while not nombre:
        nombre = pedir("  Nombre completo de la enfermera", "").strip()
    ced = ""
    while not ced.isdigit():
        ced = limpiar(pedir("  Cédula (solo números)", ""))
    return {"nombre_prof": nombre.upper(), "cedula": ced,
            "tipo_id_prof": "cc__c_dula_de_cuidadan_a",
            "equipo": pedir("  Número del equipo de salud", ID_EQUIPO),
            "territorio": pedir("  Territorio", TERRITORIO),
            "ficha": pedir("  Código de la ficha", CODIGO_FICHA)}


def main():
    from selenium import webdriver
    from selenium.webdriver.chrome.service import Service
    from webdriver_manager.chrome import ChromeDriverManager
    print("=" * 55 + "\n  BOT CARACTERIZACIÓN RURAL — RES 272 DE 2026\n" + "=" * 55)
    try:
        familias, veredas = leer_familias(ARCHIVO_EXCEL, HOJA_NOMBRE)
    except Exception as e:
        print(f"\n✘ Error Excel: {e}"); return
    print(f"\n✔ {len(familias)} familia(s):")
    sin_perfil = set()
    for i, f in enumerate(familias, 1):
        m = norm(f[0]["micro"])
        if m not in veredas:
            sin_perfil.add(f[0]["micro"])
        print(f"  {i:03d}: {' | '.join(p['nombre'] for p in f)}")
    if sin_perfil:
        print(f"\n⚠ Sin fila en la hoja '{HOJA_VEREDAS}' (se usa perfil "
              f"{PERFIL_VEREDA_DEFECTO}): {sorted(sin_perfil)}")
    for m, pf in veredas.items():
        if pf not in PERFILES:
            print(f"✘ Perfil '{pf}' de la vereda {m} no existe. Perfiles: {list(PERFILES)}"); return

    run = pedir_datos_corrida()
    while True:
        r = input("\n¿En qué número de familia/hogar empezamos? [1]: ").strip() or "1"
        if r.isdigit() and int(r) >= 1:
            ini = int(r); break
    print(f"✔ Consecutivos {str(ini).zfill(3)} a {str(ini + len(familias) - 1).zfill(3)}")
    url = pedir("\nLink del formulario", URL_FORMULARIO)
    input("\nENTER para comenzar (Ctrl+C cancela)...")

    opts = webdriver.ChromeOptions()
    opts.add_argument("--start-maximized"); opts.add_argument("--disable-notifications")
    driver = webdriver.Chrome(service=Service(ChromeDriverManager().install()), options=opts)
    bot, ok, fallo, avisos = Bot(driver), [], None, set()
    try:
        driver.get(url); time.sleep(6)
        bot.descartar_borrador()
        bot.leer_mapa()
        faltan = [q for q, _, _ in usadas_en_codigo() if q not in bot.mapa]
        if faltan:
            raise CampoError(f"El formulario cargado no tiene las preguntas: {sorted(set(faltan))}. "
                             f"¿Es el formulario rural correcto?")
        with open(LOG_CSV, "a", newline="", encoding="utf-8") as lg:
            w = csv.writer(lg)
            for i, fam in enumerate(familias):
                cod = str(ini + i).zfill(3)
                perfil = veredas.get(norm(fam[0]["micro"]), PERFIL_VEREDA_DEFECTO)
                try:
                    llenar(bot, fam, ini + i, perfil, run, avisos)
                    ok.append((cod, fam[0]["nombre"]))
                    w.writerow([datetime.now().isoformat(timespec="seconds"), cod,
                                fam[0]["nombre"], "ENVIADA"]); lg.flush()
                except Exception as e:
                    print(f"\n  ✘✘✘ {e}")
                    fallo = (cod, fam[0]["nombre"], str(e))
                    w.writerow([datetime.now().isoformat(timespec="seconds"), cod,
                                fam[0]["nombre"], "ERROR: " + str(e)])
                    break
    finally:
        if fallo:
            print("\n(El navegador queda abierto para revisar el registro con error.)")
        else:
            driver.quit()
    print("\n" + "=" * 55 + "\n  INFORME FINAL\n" + "=" * 55)
    print(f"  ✔ Familias enviadas: {len(ok)}")
    for c, n in ok[-5:]:
        print(f"    • {c}: {n}")
    if fallo:
        print(f"\n  ⛔ DETENIDO en familia {fallo[0]} ({fallo[1]})\n     Motivo: {fallo[2]}")
        print(f"  ➜ NO se envió: reintenta desde el consecutivo {fallo[0]}.")
    else:
        print("\n  ✔ Todas las familias se cargaron.")
    if avisos:
        print("\n  Supuestos aplicados en esta corrida:")
        for a in sorted(avisos):
            print(f"   - {a}")
    print("=" * 55)


if __name__ == "__main__":
    if len(sys.argv) >= 3 and sys.argv[1] == "--validar":
        sys.exit(0 if validar_html(sys.argv[2]) else 1)
    main()
